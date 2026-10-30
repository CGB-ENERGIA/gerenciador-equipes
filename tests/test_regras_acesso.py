"""Regras de acesso e de troca pelas rotas, como uma pessoa usaria: um SQLite em
memória com chaves estrangeiras ligadas (para o ON DELETE CASCADE valer como no
Postgres). Não toca no banco de verdade."""

import io

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app as modulo_app
import auth
import indiretos
from database.base import Base
from database.models import (
    AlocacaoIndireto,
    Colaborador,
    ComposicaoEquipe,
    Equipe,
    MembroEquipe,
    NivelPermissao,
    SetorGerencia,
    Usuario,
    VinculoUsuario,
)


@pytest.fixture
def Sessao(monkeypatch):
    motor = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    event.listen(motor, "connect", lambda con, _: con.execute("PRAGMA foreign_keys=ON"))
    Base.metadata.create_all(motor)
    fabrica = sessionmaker(autocommit=False, autoflush=False, bind=motor)
    monkeypatch.setattr(modulo_app, "SessionLocal", fabrica)
    monkeypatch.setattr(auth, "SessionLocal", fabrica)
    monkeypatch.setattr(indiretos, "SessionLocal", fabrica)
    monkeypatch.setitem(modulo_app.app.config, "SECRET_KEY", "teste")

    s = fabrica()
    bacabal = Equipe(BASE="BACABAL", PREFIXO="BCB-01")
    pedreiras = Equipe(BASE="PEDREIRAS", PREFIXO="PED-01")
    s.add_all([bacabal, pedreiras])
    s.flush()
    s.add_all([
        ComposicaoEquipe(id=1, equipe_id=bacabal.id, FUNÇÃO_ER="ELETRICISTA", ESTRUTURA="CONSTRUÇÃO"),
        ComposicaoEquipe(id=2, equipe_id=bacabal.id, FUNÇÃO_ER="ELETRICISTA", ESTRUTURA="CONSTRUÇÃO"),
        ComposicaoEquipe(id=3, equipe_id=pedreiras.id, FUNÇÃO_ER="ELETRICISTA", ESTRUTURA="CONSTRUÇÃO"),
        ComposicaoEquipe(id=4, equipe_id=pedreiras.id, FUNÇÃO_ER="ELETRICISTA", ESTRUTURA="CONSTRUÇÃO"),
        Colaborador(CHAPA="10", NOME="ANA DA MINHA EQUIPE", TIPO_FUNÇÃO="DIRETO", AFASTADO=False),
        Colaborador(CHAPA="30", NOME="BIA DE OUTRA BASE", TIPO_FUNÇÃO="DIRETO", AFASTADO=False),
        Colaborador(CHAPA="40", NOME="CAIO LIVRE", TIPO_FUNÇÃO="DIRETO", AFASTADO=False),
        Colaborador(
            CHAPA="50", NOME="DITO AFASTADO", TIPO_FUNÇÃO="DIRETO", AFASTADO=True,
            JUSTIFICATIVA_AFASTAMENTO="LICENÇA", AFASTADO_POR="alguem",
        ),
        Colaborador(CHAPA="60", NOME="EVA INDIRETA", TIPO_FUNÇÃO="INDIRETO", AFASTADO=False),
        MembroEquipe(composicao_id=1, CHAPA="10"),
        MembroEquipe(composicao_id=3, CHAPA="30"),
        Usuario(id=1, USUARIO="admin", NOME="ADMIN", SENHA_HASH="x", NIVEL="ADMINISTRADOR", ATIVO=True),
        Usuario(id=2, USUARIO="sup", NOME="SUP", SENHA_HASH="x", NIVEL="SUPERVISOR", ATIVO=True),
        Usuario(id=3, USUARIO="semvinculo", NOME="SEM", SENHA_HASH="x", NIVEL="SUPERVISOR", ATIVO=True),
        Usuario(id=4, USUARIO="gerente", NOME="GER", SENHA_HASH="x", NIVEL="GERENTE", ATIVO=True),
        Usuario(id=5, USUARIO="outroadmin", NOME="ADMIN 2", SENHA_HASH="y", NIVEL="ADMINISTRADOR", ATIVO=True),
        VinculoUsuario(usuario_id=2, TIPO="BASE", VALOR="BACABAL"),
        VinculoUsuario(usuario_id=2, TIPO="TIPO_EQUIPE", VALOR="CONSTRUÇÃO"),
    ])
    s.commit()
    s.close()
    return fabrica


def cliente(usuario_id, senha_hash="x"):
    c = modulo_app.app.test_client()
    with c.session_transaction() as sessao_http:
        sessao_http[auth.CHAVE_SESSAO] = usuario_id
        sessao_http[auth.CHAVE_MARCA_SENHA] = auth.marca_da_senha(senha_hash)
    return c


def ler(Sessao, consulta):
    s = Sessao()
    try:
        return consulta(s)
    finally:
        s.close()


def ocupante(Sessao, vaga):
    membro = ler(Sessao, lambda s: s.query(MembroEquipe).filter_by(composicao_id=vaga).first())
    return membro.CHAPA if membro else None


def colaborador(Sessao, chapa):
    return ler(Sessao, lambda s: s.query(Colaborador).filter_by(CHAPA=chapa).first())


def so_consulta(Sessao, nivel="SUPERVISOR"):
    s = Sessao()
    s.add(NivelPermissao(NIVEL=nivel, PERMISSAO=auth.VER_EQUIPES))
    s.commit()
    s.close()


# ------------------------------------------------------------
# Decisão A: traz ou substitui para as suas equipes, não edita as dos outros
# ------------------------------------------------------------

def test_supervisor_traz_colaborador_de_outra_base_com_confirmacao(Sessao):
    sup = cliente(2)
    r = sup.post("/api/equipes/alocar", json={"composicao_id": 2, "chapa": "30"})
    assert r.status_code == 409 and r.get_json()["conflito"]
    assert ocupante(Sessao, 3) == "30"

    r = sup.post("/api/equipes/alocar", json={"composicao_id": 2, "chapa": "30", "confirmar_transferencia": True})
    assert r.status_code == 200
    assert ocupante(Sessao, 2) == "30" and ocupante(Sessao, 3) is None


def test_supervisor_substitui_na_sua_vaga_por_colaborador_de_outra_base(Sessao):
    r = cliente(2).post("/api/equipes/editar-alocacao", json={
        "composicao_id": 1, "chapa": "30", "confirmar_transferencia": True,
    })
    assert r.status_code == 200
    assert ocupante(Sessao, 1) == "30" and ocupante(Sessao, 3) is None


@pytest.mark.parametrize(
    ("rota", "corpo"),
    [
        ("/api/equipes/alocar", {"composicao_id": 4, "chapa": "40"}),
        ("/api/equipes/editar-alocacao", {"composicao_id": 3, "chapa": "40"}),
        ("/api/equipes/remover", {"composicao_id": 3}),
        ("/api/colaboradores/afastar", {"chapa": "30", "justificativa": "X"}),
        ("/api/colaboradores/observacao", {"chapa": "30", "observacao": "X"}),
        ("/api/colaboradores/tipo-funcao", {"chapa": "30", "tipo_funcao": "INDIRETO"}),
    ],
)
def test_supervisor_nao_edita_outra_base(Sessao, rota, corpo):
    assert cliente(2).post(rota, json=corpo).status_code == 403
    assert ocupante(Sessao, 3) == "30" and ocupante(Sessao, 4) is None
    pessoa = colaborador(Sessao, "30")
    assert not pessoa.AFASTADO and not pessoa.OBSERVACAO and pessoa.TIPO_FUNÇÃO == "DIRETO"


def test_supervisor_nao_esvazia_equipe_de_outra_base(Sessao):
    equipe = ler(Sessao, lambda s: s.query(Equipe).filter_by(PREFIXO="PED-01").first().id)
    assert cliente(2).delete(f"/api/equipes/{equipe}/membros").status_code == 403
    assert ocupante(Sessao, 3) == "30"


def test_supervisor_opera_a_propria_base(Sessao):
    sup = cliente(2)
    assert sup.post("/api/colaboradores/observacao", json={"chapa": "10", "observacao": "X"}).status_code == 200
    assert sup.post("/api/equipes/remover", json={"composicao_id": 1}).status_code == 200


def test_folguista_extra_de_outra_base_exige_vinculo(Sessao):
    s = Sessao()
    folguista = Equipe(BASE="PEDREIRAS", PREFIXO="FOLGUISTA")
    s.add(folguista)
    s.flush()
    vaga = ComposicaoEquipe(
        equipe_id=folguista.id, FUNÇÃO_ER="ELETRICISTA", ESTRUTURA="CONSTRUÇÃO", ORIGEM="EXTRA"
    )
    s.add(vaga)
    s.flush()
    s.add(MembroEquipe(composicao_id=vaga.id, CHAPA="40"))
    s.commit()
    equipe_id, vaga_id = folguista.id, vaga.id
    s.close()

    sup = cliente(2)
    r = sup.post(f"/api/equipes/{equipe_id}/folguista-extra", json={
        "funcao_er": "ELETRICISTA", "chapa": "50", "setor": modulo_app.SETORES_NEGOCIO[0],
    })
    assert r.status_code == 403
    assert sup.post("/api/equipes/remover", json={"composicao_id": vaga_id}).status_code == 403
    assert ocupante(Sessao, vaga_id) == "40"


def test_administrador_continua_sem_restricao(Sessao):
    assert cliente(1).post("/api/equipes/remover", json={"composicao_id": 3}).status_code == 200


# ------------------------------------------------------------
# Troca: afastado e indireto
# ------------------------------------------------------------

def test_troca_retira_o_afastamento(Sessao):
    r = cliente(2).post("/api/equipes/editar-alocacao", json={"composicao_id": 1, "chapa": "50"})
    assert r.status_code == 200
    assert ocupante(Sessao, 1) == "50"
    pessoa = colaborador(Sessao, "50")
    assert not pessoa.AFASTADO
    assert pessoa.JUSTIFICATIVA_AFASTAMENTO is None and pessoa.AFASTADO_POR is None


def test_troca_recusa_indireto(Sessao):
    r = cliente(2).post("/api/equipes/editar-alocacao", json={"composicao_id": 1, "chapa": "60"})
    assert r.status_code == 400
    assert ocupante(Sessao, 1) == "10"


def planilha_alocacoes(vaga, chapa):
    arquivo = io.BytesIO()
    pd.DataFrame([{"COMPOSICAO_ID": vaga, "CHAPA": chapa}]).to_excel(
        arquivo, index=False, sheet_name="Alocações"
    )
    arquivo.seek(0)
    return {"arquivo": (arquivo, "alocacoes.xlsx")}


def test_planilha_sinaliza_e_retira_o_afastamento(Sessao):
    sup = cliente(2)
    previa = sup.post("/api/alocacoes/planilha/previa", data=planilha_alocacoes(2, 50)).get_json()
    assert previa["alocar"][0]["afastado"] is True

    assert sup.post("/api/alocacoes/planilha/aplicar", data=planilha_alocacoes(2, 50)).status_code == 200
    assert ocupante(Sessao, 2) == "50" and not colaborador(Sessao, "50").AFASTADO


def test_planilha_recusa_indireto(Sessao):
    previa = cliente(2).post("/api/alocacoes/planilha/previa", data=planilha_alocacoes(2, 60)).get_json()
    assert previa["erros"] and not previa["alocar"]


# ------------------------------------------------------------
# Afastar exige permissão e vínculo
# ------------------------------------------------------------

def test_nivel_de_consulta_nao_afasta_nem_reativa(Sessao):
    so_consulta(Sessao)
    sup = cliente(2)
    for chapa in ("10", "40"):
        r = sup.post("/api/colaboradores/afastar", json={"chapa": chapa, "justificativa": "X"})
        assert r.status_code == 403
    assert sup.post("/api/colaboradores/reativar", json={"chapa": "50"}).status_code == 403
    assert ocupante(Sessao, 1) == "10"
    assert colaborador(Sessao, "50").AFASTADO


def test_afastar_alocado_exige_a_permissao_de_remover(Sessao):
    s = Sessao()
    s.add_all([
        NivelPermissao(NIVEL="SUPERVISOR", PERMISSAO=auth.VER_EQUIPES),
        NivelPermissao(NIVEL="SUPERVISOR", PERMISSAO=auth.ALOCAR),
    ])
    s.commit()
    s.close()
    sup = cliente(2)
    assert sup.post("/api/colaboradores/afastar", json={"chapa": "10", "justificativa": "X"}).status_code == 403
    assert sup.post("/api/colaboradores/afastar", json={"chapa": "40", "justificativa": "X"}).status_code == 200


def test_conta_sem_vinculo_nao_afasta_ninguem(Sessao):
    r = cliente(3).post("/api/colaboradores/afastar", json={"chapa": "40", "justificativa": "X"})
    assert r.status_code == 403
    assert not colaborador(Sessao, "40").AFASTADO


def test_supervisor_afasta_livre_e_alocado_da_propria_base(Sessao):
    sup = cliente(2)
    assert sup.post("/api/colaboradores/afastar", json={"chapa": "40", "justificativa": "X"}).status_code == 200
    assert sup.post("/api/colaboradores/afastar", json={"chapa": "10", "justificativa": "X"}).status_code == 200
    assert ocupante(Sessao, 1) is None


# ------------------------------------------------------------
# Conta sem vínculo não acessa dados
# ------------------------------------------------------------

@pytest.mark.parametrize(
    "rota",
    [
        "/api/resumo",
        "/api/pessoas-nao-alocadas",
        "/api/pessoas-afastadas",
        "/api/equipes",
        "/api/alocacoes/planilha",
        "/api/opcoes-alocacao",
        "/api/colaboradores",
        "/api/indiretos",
    ],
)
def test_conta_sem_vinculo_nao_le_dados(Sessao, rota):
    assert cliente(3).get(rota).status_code == 403


def test_sessao_sem_vinculo_avisa_e_nao_entrega_permissao_de_dados(Sessao):
    dados = cliente(3).get("/api/sessao").get_json()["usuario"]
    assert dados["sem_vinculo"] is True and dados["permissoes"] == []


def test_vinculo_proprio_e_administrador_continuam_lendo(Sessao):
    for usuario_id in (1, 2):
        assert cliente(usuario_id).get("/api/colaboradores").status_code == 200
    assert cliente(2).get("/api/sessao").get_json()["usuario"]["sem_vinculo"] is False


def test_gerente_com_vinculo_herdado_de_subordinado_continua_lendo(Sessao):
    s = Sessao()
    s.add_all([
        Usuario(id=6, USUARIO="coord", NOME="COORD", SENHA_HASH="x", NIVEL="COORDENADOR", ATIVO=True, RESPONSAVEL_ID=4),
        Usuario(id=7, USUARIO="sup2", NOME="SUP 2", SENHA_HASH="x", NIVEL="SUPERVISOR", ATIVO=True, RESPONSAVEL_ID=6),
    ])
    s.flush()
    s.add(VinculoUsuario(usuario_id=7, TIPO="BASE", VALOR="PEDREIRAS"))
    s.commit()
    s.close()
    assert cliente(4).get("/api/colaboradores").status_code == 200
    assert cliente(6).get("/api/colaboradores").status_code == 200


def test_gerenciar_usuarios_continua_sem_vinculo_mas_o_resto_nao(Sessao, gerente_de_usuarios):
    assert gerente_de_usuarios.get("/api/usuarios").status_code == 200
    assert gerente_de_usuarios.get("/api/colaboradores").status_code == 403


# ------------------------------------------------------------
# Troca de chapa mantém os vínculos
# ------------------------------------------------------------

def indireto_com_equipe(Sessao):
    s = Sessao()
    s.add_all([
        Colaborador(CHAPA="70", NOME="FABIO CHEFE", TIPO_FUNÇÃO="INDIRETO", AFASTADO=False),
        Colaborador(CHAPA="71", NOME="GIL SUBORDINADO", TIPO_FUNÇÃO="INDIRETO", AFASTADO=False),
    ])
    s.flush()
    s.add_all([
        AlocacaoIndireto(CHAPA="70", SETOR="FROTA", BASE="BCB", SUBSETOR="OFICINA"),
        AlocacaoIndireto(CHAPA="71", SETOR="FROTA", BASE="BCB", RESPONSAVEL_CHAPA="70"),
        SetorGerencia(SETOR="FROTA", RESPONSAVEL_CHAPA="70", GERENTE_CHAPA="71"),
        SetorGerencia(SETOR="OBRAS", RESPONSAVEL_CHAPA="", GERENTE_CHAPA="70"),
    ])
    s.commit()
    s.close()


def corpo_cadastro(chapa, nome):
    return {
        "chapa": chapa, "nome": nome, "tipo_funcao": "INDIRETO",
        "rateios": [{"rateio": "R1", "grpccusto": "1"}],
    }


def test_trocar_so_a_chapa_mantem_os_vinculos(Sessao):
    indireto_com_equipe(Sessao)
    r = cliente(1).put("/api/colaboradores/cadastro/70", json=corpo_cadastro("7000", "Fabio Chefe"))
    assert r.status_code == 200, r.get_json()

    s = Sessao()
    assert not s.query(Colaborador).filter_by(CHAPA="70").first()
    propria = s.query(AlocacaoIndireto).filter_by(CHAPA="7000").one()
    assert (propria.SETOR, propria.SUBSETOR) == ("FROTA", "OFICINA")
    assert s.query(AlocacaoIndireto).filter_by(CHAPA="71").one().RESPONSAVEL_CHAPA == "7000"
    assert s.query(SetorGerencia).filter_by(SETOR="FROTA").one().RESPONSAVEL_CHAPA == "7000"
    assert s.query(SetorGerencia).filter_by(SETOR="OBRAS").one().GERENTE_CHAPA == "7000"
    s.close()


def test_trocar_chapa_de_alocado_mantem_a_vaga(Sessao):
    r = cliente(1).put("/api/colaboradores/cadastro/10", json={
        **corpo_cadastro("1000", "ANA DA MINHA EQUIPE"), "tipo_funcao": "DIRETO",
    })
    assert r.status_code == 200, r.get_json()
    assert ocupante(Sessao, 1) == "1000"


def test_trocar_chapa_e_nome_juntos_e_recusado_sem_perder_nada(Sessao):
    indireto_com_equipe(Sessao)
    r = cliente(1).put("/api/colaboradores/cadastro/70", json=corpo_cadastro("7000", "OUTRA PESSOA"))
    assert r.status_code == 409

    s = Sessao()
    assert s.query(AlocacaoIndireto).filter_by(CHAPA="70").one().SETOR == "FROTA"
    assert not s.query(Colaborador).filter_by(CHAPA="7000").first()
    s.close()


# ------------------------------------------------------------
# Ciclo na cadeia de responsáveis
# ------------------------------------------------------------

def alocar_indireto_corpo(chapa, responsavel):
    return {"chapa": chapa, "setor": "FROTA", "base": "BCB", "responsavel_chapa": responsavel}


def test_formulario_e_troca_de_responsavel_recusam_ciclo(Sessao):
    indireto_com_equipe(Sessao)  # 71 responde a 70
    admin = cliente(1)
    assert admin.post("/api/indiretos/alocar", json=alocar_indireto_corpo("70", "71")).status_code == 400
    r = admin.post("/api/indiretos/responsavel", json={"chapas": ["70"], "responsavel_chapa": "71"})
    assert r.status_code == 400
    assert ler(Sessao, lambda s: s.query(AlocacaoIndireto).filter_by(CHAPA="70").one().RESPONSAVEL_CHAPA) is None


def test_chefe_de_setor_continua_respondendo_a_si_mesmo(Sessao):
    indireto_com_equipe(Sessao)
    r = cliente(1).post("/api/indiretos/alocar", json=alocar_indireto_corpo("70", "70"))
    assert r.status_code == 200, r.get_json()


def planilha_indiretos(linhas):
    colunas = ["SETOR", "BASE", "CHAPA RESPONSÁVEL", "CHAPA INDIRETO"]
    arquivo = io.BytesIO()
    pd.DataFrame(linhas, columns=colunas).to_excel(arquivo, index=False, sheet_name="Alocação")
    arquivo.seek(0)
    return arquivo


def test_planilha_de_indiretos_recusa_ciclo_entre_linhas(Sessao):
    indireto_com_equipe(Sessao)
    s = Sessao()
    plano = indiretos.analisar_planilha_indiretos(planilha_indiretos([
        ["FROTA", "BCB", 71, 70],
        ["FROTA", "BCB", 70, 71],
    ]), s)
    s.close()
    assert len(plano["erros"]) >= 1


# ------------------------------------------------------------
# Excluir colaborador com subordinados
# ------------------------------------------------------------

def test_excluir_responsavel_pede_confirmacao_e_nao_deixa_chapa_solta(Sessao):
    indireto_com_equipe(Sessao)
    admin = cliente(1)

    r = admin.delete("/api/colaboradores/cadastro/70")
    assert r.status_code == 409
    assert r.get_json()["subordinados"] == 1 and r.get_json()["setores"] == 2
    assert colaborador(Sessao, "70")

    assert admin.delete("/api/colaboradores/cadastro/70?confirmar_subordinados=1").status_code == 200
    assert not colaborador(Sessao, "70")
    assert ler(Sessao, lambda s: s.query(AlocacaoIndireto).filter_by(CHAPA="71").one().RESPONSAVEL_CHAPA) is None
    assert ler(Sessao, lambda s: s.query(SetorGerencia).count()) == 0


def test_excluir_em_massa_pede_confirmacao_e_ignora_quem_tambem_sai(Sessao):
    indireto_com_equipe(Sessao)
    admin = cliente(1)

    r = admin.post("/api/colaboradores/cadastro/excluir-massa", json={"chapas": ["70"]})
    assert r.status_code == 409 and r.get_json()["subordinados"] == 1

    r = admin.post("/api/colaboradores/cadastro/excluir-massa", json={"chapas": ["70", "71"]})
    assert r.status_code == 409 and r.get_json()["subordinados"] == 0

    r = admin.post("/api/colaboradores/cadastro/excluir-massa", json={
        "chapas": ["70", "71"], "confirmar_subordinados": True,
    })
    assert r.status_code == 200


def test_excluir_sem_dependentes_nao_pede_confirmacao(Sessao):
    assert cliente(1).delete("/api/colaboradores/cadastro/40").status_code == 200


# ------------------------------------------------------------
# Texto com "=" nas exportações
# ------------------------------------------------------------

def celulas_da_resposta(resposta):
    livro = load_workbook(io.BytesIO(resposta.data))
    return {
        (aba.title, celula.value): celula
        for aba in livro.worksheets for linha in aba.iter_rows() for celula in linha
    }


def test_exportacao_grava_texto_com_igual_como_texto(Sessao):
    s = Sessao()
    s.query(Colaborador).filter_by(CHAPA="10").update({"OBSERVACAO": "=1+1"})
    s.commit()
    s.close()

    r = cliente(1).post("/api/exportar-pessoas", json={"alocadas": [{"chapa": "10"}]})
    assert r.status_code == 200, r.get_json()
    celula = celulas_da_resposta(r)[("Colaboradores", "=1+1")]
    assert celula.data_type == "s"


def test_planilha_de_indiretos_protege_texto_e_mantem_as_formulas_dela(Sessao):
    indireto_com_equipe(Sessao)
    s = Sessao()
    s.query(AlocacaoIndireto).filter_by(CHAPA="70").update({"OBSERVACAO": "=2+2"})
    s.commit()
    s.close()

    r = cliente(1).get("/api/indiretos/planilha")
    assert r.status_code == 200
    aba = load_workbook(io.BytesIO(r.data))["Alocação"]
    linha = next(n for n in range(2, aba.max_row + 1) if aba.cell(row=n, column=5).value == "70")
    assert aba.cell(row=linha, column=8).data_type == "s"
    assert aba.cell(row=linha, column=6).data_type == "f"


# ------------------------------------------------------------
# Sessão depois de trocar a senha
# ------------------------------------------------------------

def test_trocar_a_senha_derruba_as_outras_sessoes_mas_nao_a_atual(Sessao):
    hash_atual = auth.gerar_hash_senha("senha-atual")
    s = Sessao()
    s.query(Usuario).filter_by(id=2).update({"SENHA_HASH": hash_atual})
    s.commit()
    s.close()

    outro_computador = cliente(2, hash_atual)
    este_computador = cliente(2, hash_atual)
    assert outro_computador.get("/api/equipes").status_code == 200

    r = este_computador.put("/api/minha-senha", json={"senha_atual": "senha-atual", "senha_nova": "senha-nova"})
    assert r.status_code == 200, r.get_json()

    assert este_computador.get("/api/equipes").status_code == 200
    assert outro_computador.get("/api/equipes").status_code == 401


def test_administrador_troca_a_senha_de_outro_e_derruba_a_sessao_dele(Sessao):
    sup = cliente(2)
    assert sup.get("/api/equipes").status_code == 200

    admin = cliente(1)
    r = admin.put("/api/usuarios/2", json=corpo_usuario("sup", "SUP", "SUPERVISOR", senha="senha-nova"))
    assert r.status_code == 200, r.get_json()
    assert sup.get("/api/equipes").status_code == 401
    assert admin.get("/api/equipes").status_code == 200


def test_administrador_troca_a_propria_senha_e_continua_logado(Sessao):
    admin = cliente(1)
    r = admin.put("/api/usuarios/1", json=corpo_usuario("admin", "ADMIN", "ADMINISTRADOR", senha="senha-nova"))
    assert r.status_code == 200, r.get_json()
    assert admin.get("/api/equipes").status_code == 200


def test_cookie_sem_marca_da_senha_nao_vale(Sessao):
    c = modulo_app.app.test_client()
    with c.session_transaction() as sessao_http:
        sessao_http[auth.CHAVE_SESSAO] = 1
    assert c.get("/api/equipes").status_code == 401


# ------------------------------------------------------------
# Conta de Administrador e níveis: só Administrador
# ------------------------------------------------------------

@pytest.fixture
def gerente_de_usuarios(Sessao):
    s = Sessao()
    s.add_all([
        NivelPermissao(NIVEL="GERENTE", PERMISSAO=auth.VER_EQUIPES),
        NivelPermissao(NIVEL="GERENTE", PERMISSAO=auth.GERENCIAR_USUARIOS),
    ])
    s.commit()
    s.close()
    return cliente(4)


def corpo_usuario(login, nome, nivel, **extra):
    return {"usuario": login, "nome": nome, "nivel": nivel, "ativo": True, **extra}


def test_gerente_de_usuarios_nao_altera_administrador(Sessao, gerente_de_usuarios):
    r = gerente_de_usuarios.put("/api/usuarios/5", json=corpo_usuario(
        "outroadmin", "ADMIN 2", "ADMINISTRADOR", senha="senha-nova",
    ))
    assert r.status_code == 403
    assert ler(Sessao, lambda s: s.query(Usuario).filter_by(id=5).one().SENHA_HASH) == "y"

    r = gerente_de_usuarios.put("/api/usuarios/5", json=corpo_usuario("outroadmin", "ADMIN 2", "SUPERVISOR"))
    assert r.status_code == 403
    assert gerente_de_usuarios.delete("/api/usuarios/5").status_code == 403
    assert gerente_de_usuarios.post("/api/usuarios/5/desbloquear").status_code == 403
    assert ler(Sessao, lambda s: s.query(Usuario).filter_by(id=5).one().NIVEL) == "ADMINISTRADOR"


def test_gerente_de_usuarios_nao_edita_niveis(Sessao, gerente_de_usuarios):
    r = gerente_de_usuarios.put("/api/niveis/GERENTE/permissoes", json={"permissoes": list(auth.TODAS_PERMISSOES)})
    assert r.status_code == 403
    assert auth.GERENCIAR_VAGAS not in ler(Sessao, lambda s: auth.permissoes_do_nivel("GERENTE", s))


def test_gerente_de_usuarios_nao_altera_administrador_pela_planilha(Sessao, gerente_de_usuarios):
    arquivo = io.BytesIO()
    pd.DataFrame([{"USUARIO": "outroadmin", "NOME": "ADMIN 2", "NIVEL": "Administrador", "SENHA": "senha-nova"}]).to_excel(
        arquivo, index=False
    )
    arquivo.seek(0)
    r = gerente_de_usuarios.post("/api/usuarios/planilha/aplicar", data={"arquivo": (arquivo, "u.xlsx")})
    assert r.status_code == 400
    assert ler(Sessao, lambda s: s.query(Usuario).filter_by(id=5).one().SENHA_HASH) == "y"


def test_gerente_de_usuarios_altera_quem_nao_e_administrador(Sessao, gerente_de_usuarios):
    r = gerente_de_usuarios.put("/api/usuarios/3", json=corpo_usuario("semvinculo", "NOVO NOME", "SUPERVISOR"))
    assert r.status_code == 200


def test_administrador_altera_administrador_e_niveis(Sessao):
    admin = cliente(1)
    assert admin.put("/api/usuarios/5", json=corpo_usuario("outroadmin", "RENOMEADO", "ADMINISTRADOR")).status_code == 200
    r = admin.put("/api/niveis/GERENTE/permissoes", json={"permissoes": [auth.VER_EQUIPES]})
    assert r.status_code == 200
    assert admin.delete("/api/usuarios/5").status_code == 200
