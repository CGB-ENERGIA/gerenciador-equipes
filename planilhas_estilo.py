"""Identidade visual das planilhas Excel que o sistema gera.

Um lugar só para o visual de toda planilha baixada (equipes, alocações,
colaboradores, usuários, exportação do Resumo): cabeçalho bordô da CGB,
Arial, bordas suaves, colunas de preenchimento em amarelo e de consulta em
cinza, listas suspensas e uma aba "Instruções" no fim.

Cuidado com as planilhas de IMPORTAÇÃO: o sistema lê de volta a PRIMEIRA aba
(ou a aba pelo nome), com o cabeçalho na linha 1. Por isso aqui:
- nada é inserido acima do cabeçalho nem como coluna extra;
- nenhuma fórmula vai nas abas de dados (o openpyxl não grava o valor
  calculado, e o pandas leria vazio);
- o amarelo das colunas de preenchimento além das linhas com dados é feito
  por formatação condicional, que não cria células — linhas "vazias mas
  formatadas" virariam linhas fantasma na leitura;
- a aba de instruções entra sempre por ÚLTIMO.
"""

from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BORDO = "711424"
FONTE = "Arial"

FONTE_NORMAL = Font(name=FONTE, size=10)
FONTE_CABECALHO = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
FONTE_EXEMPLO = Font(name=FONTE, size=10, italic=True, color="6D5F63")
FUNDO_CABECALHO = PatternFill("solid", fgColor=BORDO)
FUNDO_PREENCHER = PatternFill("solid", fgColor="FFF9DB")   # amarelo claro
FUNDO_CONSULTA = PatternFill("solid", fgColor="F2EEEF")    # cinza-rosado
FUNDO_ZEBRA = PatternFill("solid", fgColor="FAF7F8")
FUNDO_EXEMPLO = PatternFill("solid", fgColor="FAF7F8")
FONTE_ALERTA = Font(name=FONTE, size=10, bold=True, color="8A5A00")
LINHA = Side(style="thin", color="D9C9CD")


def fundo_condicional(cor):
    """Fundo para formatação condicional. Na formatação condicional o Excel lê
    a cor do bgColor (não do fgColor das células comuns): sem as duas, a
    regra existe mas não pinta nada."""
    return PatternFill(start_color=cor, end_color=cor, fill_type="solid")
BORDA = Border(left=LINHA, right=LINHA, top=LINHA, bottom=LINHA)

# quantas linhas abaixo dos dados ganham o amarelo e as listas suspensas,
# para quem acrescenta linhas no fim (formatação condicional e validação não
# criam células, então não geram linhas fantasma na leitura)
LINHAS_EXTRAS = 500


def estilizar_aba(
    planilha,
    *,
    editaveis=(),
    consulta=False,
    validacoes=None,
    linhas_exemplo=(),
    largura_minima=12,
    largura_maxima=48,
):
    """Aplica a identidade da CGB numa aba já preenchida (cabeçalho na linha 1).

    editaveis: nomes das colunas que a pessoa preenche (fundo amarelo). Numa
        aba de importação, as demais ficam cinza (só referência).
    consulta: aba só de leitura — linhas zebradas, sem amarelo.
    validacoes: {coluna: [valores]} para lista suspensa (aceita outro valor
        com aviso, porque o sistema é quem valida de verdade).
    linhas_exemplo: números de linha (Excel) que são só exemplo — cinza itálico.
    """
    ultima_linha = max(planilha.max_row, 1)
    ultima_coluna = max(planilha.max_column, 1)
    cabecalhos = {
        planilha.cell(row=1, column=c).value: c for c in range(1, ultima_coluna + 1)
    }
    colunas_editaveis = {cabecalhos[n] for n in editaveis if n in cabecalhos}

    # ---------- cabeçalho
    for c in range(1, ultima_coluna + 1):
        celula = planilha.cell(row=1, column=c)
        celula.font = FONTE_CABECALHO
        celula.fill = FUNDO_CABECALHO
        celula.border = BORDA
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if editaveis and not consulta:
            celula.comment = Comment(
                "Preencha / edite esta coluna." if c in colunas_editaveis
                else "Só referência: não precisa mexer.",
                "Gerenciador de Equipes",
            )
    planilha.row_dimensions[1].height = 30

    # ---------- corpo
    for linha in range(2, ultima_linha + 1):
        for c in range(1, ultima_coluna + 1):
            celula = planilha.cell(row=linha, column=c)
            celula.border = BORDA
            celula.alignment = Alignment(vertical="center")
            if linha in linhas_exemplo:
                celula.font, celula.fill = FONTE_EXEMPLO, FUNDO_EXEMPLO
                continue
            celula.font = FONTE_NORMAL
            if consulta:
                if linha % 2 == 1:
                    celula.fill = FUNDO_ZEBRA
            elif editaveis:
                celula.fill = FUNDO_PREENCHER if c in colunas_editaveis else FUNDO_CONSULTA

    # amarelo também nas linhas que a pessoa acrescentar no fim
    if editaveis and not consulta:
        for c in colunas_editaveis:
            letra = get_column_letter(c)
            inicio = ultima_linha + 1
            fim = ultima_linha + LINHAS_EXTRAS
            planilha.conditional_formatting.add(
                f"{letra}{inicio}:{letra}{fim}",
                FormulaRule(formula=["TRUE"], fill=fundo_condicional("FFF9DB")),
            )

    # ---------- listas suspensas
    for coluna, valores in (validacoes or {}).items():
        if coluna not in cabecalhos or not valores:
            continue
        letra = get_column_letter(cabecalhos[coluna])
        lista = ",".join(str(v) for v in valores)
        validacao = DataValidation(
            type="list",
            formula1=f'"{lista}"',
            allow_blank=True,
            errorStyle="warning",
            errorTitle="Valor fora da lista",
            error=f"Valores esperados: {lista.replace(',', ', ')}.",
        )
        planilha.add_data_validation(validacao)
        validacao.add(f"{letra}2:{letra}{ultima_linha + LINHAS_EXTRAS}")

    # ---------- larguras, congelamento, filtro, cor da aba
    for c in range(1, ultima_coluna + 1):
        maior_dado = max(
            (len(str(planilha.cell(row=l, column=c).value or ""))
             for l in range(2, min(ultima_linha, 400) + 1)),
            default=0,
        )
        # cabeçalho em negrito e com a setinha do filtro pede mais espaço,
        # senão quebra no meio da palavra ("ENCARREGAD / O")
        cabecalho = len(str(planilha.cell(row=1, column=c).value or ""))
        largura = max(maior_dado * 1.1 + 3, cabecalho * 1.2 + 5)
        planilha.column_dimensions[get_column_letter(c)].width = max(
            largura_minima, min(largura, largura_maxima)
        )
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = f"A1:{get_column_letter(ultima_coluna)}{ultima_linha}"
    planilha.sheet_properties.tabColor = "9C8B90" if consulta else BORDO


def destacar_valor(planilha, coluna, valor):
    """Pinta de âmbar as células da coluna iguais a `valor` (ex.: AFASTADO=SIM)."""
    cabecalhos = {
        planilha.cell(row=1, column=c).value: c for c in range(1, planilha.max_column + 1)
    }
    if coluna not in cabecalhos:
        return
    letra = get_column_letter(cabecalhos[coluna])
    planilha.conditional_formatting.add(
        f"{letra}2:{letra}{max(planilha.max_row, 2)}",
        FormulaRule(
            formula=[f'{letra}2="{valor}"'], fill=fundo_condicional("FDF0CD"), font=FONTE_ALERTA
        ),
    )


def adicionar_instrucoes(livro, titulo, itens, *, legenda=True):
    """Aba "Instruções" no FIM do arquivo (a leitura usa a primeira aba).

    itens: lista de (tópico, explicação).
    """
    aba = livro.create_sheet("Instruções")
    aba.sheet_view.showGridLines = False
    aba.sheet_properties.tabColor = "1B6B45"
    aba.column_dimensions["A"].width = 3
    aba.column_dimensions["B"].width = 26
    aba.column_dimensions["C"].width = 95

    aba["B2"] = titulo
    aba["B2"].font = Font(name=FONTE, size=14, bold=True, color=BORDO)
    aba["B3"] = "Gerenciador de Equipes · CGB Energia"
    aba["B3"].font = Font(name=FONTE, size=10, bold=True, color="8D2136")

    linha = 5
    if legenda:
        for fundo, texto in (
            (FUNDO_PREENCHER, "Fundo amarelo: coluna para preencher ou editar."),
            (FUNDO_CONSULTA, "Fundo cinza: só referência, não precisa mexer."),
        ):
            aba.cell(row=linha, column=2).fill = fundo
            aba.cell(row=linha, column=2).border = BORDA
            celula = aba.cell(row=linha, column=3, value=texto)
            celula.font = FONTE_NORMAL
            linha += 1
        linha += 1

    for topico, explicacao in itens:
        a = aba.cell(row=linha, column=2, value=topico)
        a.font = Font(name=FONTE, size=10, bold=True, color=BORDO)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = aba.cell(row=linha, column=3, value=explicacao)
        b.font = FONTE_NORMAL
        b.alignment = Alignment(vertical="top", wrap_text=True)
        aba.row_dimensions[linha].height = 15 * (len(explicacao) // 90 + 1)
        linha += 1
    return aba
