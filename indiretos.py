"""Alocação do pessoal INDIRETO (página "Indiretos").

Controle nominal, sem vagas: cada indireto fica num SETOR, numa BASE (que
quem aloca escolhe, independente do cadastro) e sob um RESPONSÁVEL — por
pessoa, então no mesmo setor e base pode haver mais de um. Tabela
alocacoes_indiretos (migrations/018), uma linha por indireto.

Rotas num Blueprint próprio para não engordar o app.py. Todos que veem
equipes consultam a página (GET /api/indiretos); alterar — alocar, mover,
trocar responsável, planilha — continua só com o Administrador
(auth.exige_administrador).

A planilha segue o padrão das outras (baixar -> preencher -> prévia ->
aplicar) e, como a de usuários, NUNCA tira ninguém: quem não aparece na
planilha continua onde está. Para tirar alguém do setor, use a tela.
"""

import io
import re
from datetime import datetime, timezone

import pandas as pd
from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

import auth
import planilhas_estilo as estilo
from database.database import SessionLocal
from database.models import AlocacaoIndireto, Colaborador, SetorGerencia

bp = Blueprint("indiretos", __name__)

TIPO_INDIRETO = "INDIRETO"
# os gerentes: as bases deles dizem quais setores aparecem abaixo de cada um
# no organograma geral, sem mudar o responsável de ninguém
SETOR_GERENCIA = "GERÊNCIA"
# gerente sem base definida: a coluna BASE é obrigatória, e esta não é uma
# base de verdade (não aparece entre as opções)
BASE_GERENCIA = "GERAL"
BASES_PADRAO = ("BCB", "BDC", "ITM", "PDS", "PDT", "STI", "SPOT STI")

ABA_ALOCACAO = "Alocação"
COL_SETOR = "SETOR"
COL_BASE = "BASE"
COL_RESPONSAVEL = "CHAPA RESPONSÁVEL"
COL_INDIRETO = "CHAPA INDIRETO"
COL_OBSERVACAO = "OBSERVAÇÃO"
COL_SUBSETOR = "SUBSETOR"
COLUNAS_OBRIGATORIAS = (COL_SETOR, COL_BASE, COL_RESPONSAVEL, COL_INDIRETO)
COLUNAS_PLANILHA = (
    COL_SETOR, COL_BASE, COL_RESPONSAVEL, "NOME RESPONSÁVEL",
    COL_INDIRETO, "NOME INDIRETO", "FUNÇÃO INDIRETO", COL_OBSERVACAO, COL_SUBSETOR,
)
LINHAS_EXTRAS_PLANILHA = 600


# ============================================================
# NORMALIZAÇÃO
# ============================================================

def texto_padrao(valor):
    """Maiúsculas, sem espaço sobrando — setor e base são comparados assim
    (mesma ideia de funcao_da_vaga no app.py)."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    return " ".join(str(valor).split()).upper()


def chapa_da_celula(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    chapa = str(valor).strip()
    if re.fullmatch(r"\d+\.0", chapa):
        chapa = chapa[:-2]
    return chapa


def texto_livre(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return None
    texto = str(valor).strip()
    return texto or None


def base_do_cadastro(colaborador):
    # import tardio: app.py importa este módulo
    from app import base_da_secao
    return base_da_secao(colaborador.SEÇÃO)["codigo"] or ""


def quem_esta_logado():
    return (auth.usuario_logado() or {}).get("usuario")


# ============================================================
# LEITURA
# ============================================================

def descrever_pessoa(colaborador, alocacao):
    return {
        "chapa": colaborador.CHAPA,
        "nome": colaborador.NOME or "",
        "funcao": colaborador.FUNÇÃO or "",
        "base_cadastro": base_do_cadastro(colaborador),
        "afastado": bool(colaborador.AFASTADO),
        "alocacao": None if not alocacao else {
            "setor": alocacao.SETOR,
            "base": alocacao.BASE,
            "responsavel_chapa": alocacao.RESPONSAVEL_CHAPA,
            "observacao": alocacao.OBSERVACAO or "",
            "funcao": alocacao.FUNCAO or "",
            "assessoria": bool(alocacao.ASSESSORIA),
            "subsetor": alocacao.SUBSETOR or "",
            "bases_responsavel": bases_da_coluna(alocacao.BASES_RESPONSAVEL),
            "abaixo_de_funcao": alocacao.ABAIXO_DE_FUNCAO or "",
            "alocado_por": alocacao.ALOCADO_POR or "",
            "alocado_em": alocacao.ALOCADO_EM.isoformat() if alocacao.ALOCADO_EM else "",
        },
    }


def subsetor_efetivo(alocacoes, chapa):
    """Subsetor da pessoa; sem um próprio, herda o do bloco abaixo do qual
    está pendurada (ABAIXO_DE_FUNCAO = "CARGO|SUBSETOR") ou, senão, o do
    chefe (e assim por diante na cadeia, dentro do mesmo setor).
    "" = setor geral. alocacoes: {chapa: AlocacaoIndireto}."""
    atual = alocacoes.get(chapa)
    setor = atual.SETOR if atual else None
    vistos = set()
    while atual and atual.CHAPA not in vistos:
        if atual.SUBSETOR:
            return atual.SUBSETOR
        _, _, do_bloco = (atual.ABAIXO_DE_FUNCAO or "").partition("|")
        if do_bloco:
            return do_bloco
        vistos.add(atual.CHAPA)
        chefe = alocacoes.get(atual.RESPONSAVEL_CHAPA)
        if not chefe or chefe.SETOR != setor:
            return ""
        atual = chefe
    return ""


def bases_da_coluna(valor):
    return [b for b in (valor or "").split(",") if b]


def funcoes_do_cadastro(colaboradores):
    """Funções que já existem no cadastro — as únicas que a página deixa
    escolher para um indireto, para manter o padrão dos nomes."""
    return sorted(
        {(c.FUNÇÃO or "").strip() for c in colaboradores} - {""},
        key=str.casefold,
    )


@bp.route("/api/indiretos", methods=["GET"])
@auth.exige_permissao(auth.VER_EQUIPES)
def obter_indiretos():
    """Tudo que a página precisa: os indiretos (com a alocação de cada um),
    quem pode ser responsável (o cadastro inteiro, para pesquisar por nome ou
    chapa) e as bases."""
    session = SessionLocal()
    try:
        alocacoes = {a.CHAPA: a for a in session.query(AlocacaoIndireto).all()}
        colaboradores = session.query(Colaborador).order_by(Colaborador.NOME).all()

        pessoas = [
            descrever_pessoa(c, alocacoes.get(c.CHAPA))
            for c in colaboradores
            if (c.TIPO_FUNÇÃO or "").upper() == TIPO_INDIRETO
        ]
        responsaveis = [
            {"chapa": c.CHAPA, "nome": c.NOME or "", "funcao": c.FUNÇÃO or ""}
            for c in colaboradores
        ]
        bases = list(BASES_PADRAO) + sorted(
            {a.BASE for a in alocacoes.values()} - set(BASES_PADRAO) - {BASE_GERENCIA}
        )
        return jsonify({
            "pessoas": pessoas,
            "responsaveis": responsaveis,
            "bases": bases,
            "funcoes": funcoes_do_cadastro(colaboradores),
            "setores_gerencia": [
                {
                    "setor": v.SETOR,
                    "responsavel_chapa": v.RESPONSAVEL_CHAPA or "",
                    "gerente_chapa": v.GERENTE_CHAPA,
                    "apoio": bool(v.APOIO),
                }
                for v in session.query(SetorGerencia).all()
            ],
        })
    except Exception as erro:
        print(f"[ERRO] obter_indiretos: {erro}")
        return jsonify({"erro": "Não foi possível carregar os indiretos."}), 500
    finally:
        session.close()


# ============================================================
# AÇÕES DA TELA
# ============================================================

def _corpo():
    dados = request.get_json(silent=True)
    return dados if isinstance(dados, dict) else {}


def _colaborador(session, chapa):
    return session.query(Colaborador).filter(Colaborador.CHAPA == chapa).first()


class FuncaoInvalida(ValueError):
    pass


def funcao_para_gravar(session, pessoa, funcao):
    """Valida a função escolhida na página e devolve o que vai para
    AlocacaoIndireto.FUNCAO: None quando vazia ou igual à do cadastro (não
    precisa guardar); erro se não existir no cadastro."""
    funcao = str(funcao or "").strip()
    if not funcao or funcao == (pessoa.FUNÇÃO or "").strip():
        return None
    if not session.query(Colaborador.id).filter(Colaborador.FUNÇÃO == funcao).first():
        raise FuncaoInvalida("Escolha uma função que já existe no cadastro.")
    return funcao


@bp.route("/api/indiretos/alocar", methods=["POST"])
@auth.exige_administrador
def alocar_indireto():
    """Aloca, move ou edita: grava setor, base, responsável e — se vierem —
    observação e função do indireto."""
    dados = _corpo()
    chapa = chapa_da_celula(dados.get("chapa"))
    setor = texto_padrao(dados.get("setor"))
    base = texto_padrao(dados.get("base"))
    responsavel = chapa_da_celula(dados.get("responsavel_chapa"))
    observacao = texto_livre(dados.get("observacao"))

    # gerente: sem base e sem responsável, fica acima de todos
    if setor == SETOR_GERENCIA:
        base = base or BASE_GERENCIA
        responsavel = responsavel or chapa

    if not (chapa and setor and base and responsavel):
        return jsonify({"erro": "Informe indireto, setor, base e responsável."}), 400

    session = SessionLocal()
    try:
        pessoa = _colaborador(session, chapa)
        if not pessoa or (pessoa.TIPO_FUNÇÃO or "").upper() != TIPO_INDIRETO:
            return jsonify({"erro": "Indireto não encontrado no cadastro."}), 404
        if not _colaborador(session, responsavel):
            return jsonify({"erro": "Responsável não encontrado no cadastro."}), 404
        try:
            funcao = (
                funcao_para_gravar(session, pessoa, dados.get("funcao"))
                if "funcao" in dados else None
            )
        except FuncaoInvalida as erro:
            return jsonify({"erro": str(erro)}), 400
        try:
            validar_cadeia(responsaveis_de_hoje(session), {chapa: responsavel}, aceita_topo=True)
        except OrganogramaInvalido as erro:
            return jsonify({"erro": str(erro)}), 400

        alocacao = session.query(AlocacaoIndireto).filter(AlocacaoIndireto.CHAPA == chapa).first()
        if not alocacao:
            alocacao = AlocacaoIndireto(CHAPA=chapa)
            session.add(alocacao)
        # outro responsável: deixa de ser assessoria do anterior e de ficar
        # abaixo de um bloco da equipe dele
        if alocacao.RESPONSAVEL_CHAPA != responsavel:
            alocacao.ASSESSORIA = False
            alocacao.ABAIXO_DE_FUNCAO = None
        # outro setor: as bases pelas quais respondia eram do setor anterior
        if alocacao.SETOR and alocacao.SETOR != setor:
            alocacao.BASES_RESPONSAVEL = None
        alocacao.SETOR = setor
        alocacao.BASE = base
        alocacao.RESPONSAVEL_CHAPA = responsavel
        if "observacao" in dados:
            alocacao.OBSERVACAO = observacao
        if "subsetor" in dados:
            alocacao.SUBSETOR = texto_padrao(dados.get("subsetor")) or None
        if "funcao" in dados:
            alocacao.FUNCAO = funcao
        alocacao.ALOCADO_POR = quem_esta_logado()
        alocacao.ALOCADO_EM = datetime.now(timezone.utc)
        session.commit()
        return jsonify({"sucesso": True})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] alocar_indireto: {erro}")
        return jsonify({"erro": "Não foi possível alocar o indireto."}), 500
    finally:
        session.close()


@bp.route("/api/indiretos/responsavel", methods=["POST"])
@auth.exige_administrador
def trocar_responsavel_indiretos():
    """Troca o responsável de um ou mais indiretos de uma vez (o grupo de um
    responsável dentro de uma seção, por exemplo)."""
    dados = _corpo()
    chapas = [chapa_da_celula(c) for c in (dados.get("chapas") or []) if chapa_da_celula(c)]
    responsavel = chapa_da_celula(dados.get("responsavel_chapa"))
    if not chapas or not responsavel:
        return jsonify({"erro": "Informe as pessoas e o novo responsável."}), 400

    session = SessionLocal()
    try:
        if not _colaborador(session, responsavel):
            return jsonify({"erro": "Responsável não encontrado no cadastro."}), 404
        try:
            validar_cadeia(
                responsaveis_de_hoje(session), {c: responsavel for c in chapas}, aceita_topo=True
            )
        except OrganogramaInvalido as erro:
            return jsonify({"erro": str(erro)}), 400
        alteradas = (
            session.query(AlocacaoIndireto)
            .filter(AlocacaoIndireto.CHAPA.in_(chapas))
            .update(
                {
                    AlocacaoIndireto.RESPONSAVEL_CHAPA: responsavel,
                    AlocacaoIndireto.ASSESSORIA: False,
                    AlocacaoIndireto.ABAIXO_DE_FUNCAO: None,
                    AlocacaoIndireto.ALOCADO_POR: quem_esta_logado(),
                    AlocacaoIndireto.ALOCADO_EM: datetime.now(timezone.utc),
                },
                synchronize_session=False,
            )
        )
        session.commit()
        return jsonify({"sucesso": True, "alteradas": alteradas})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] trocar_responsavel_indiretos: {erro}")
        return jsonify({"erro": "Não foi possível trocar o responsável."}), 500
    finally:
        session.close()


# ============================================================
# ORGANOGRAMA (arrastar pessoas)
# ============================================================

class OrganogramaInvalido(ValueError):
    pass


def responsaveis_de_hoje(session):
    return dict(session.query(AlocacaoIndireto.CHAPA, AlocacaoIndireto.RESPONSAVEL_CHAPA).all())


def validar_cadeia(responsavel_de, mudancas, aceita_topo=False):
    """responsavel_de: {chapa: responsavel_chapa} de hoje. Aplica as
    mudanças {chapa: novo_responsavel} e recusa se a cadeia fechar um ciclo
    (A responde a B que, subindo, responde a A).

    Responder a si mesmo é como o topo de um setor fica gravado: o formulário
    e a planilha aceitam (aceita_topo). Arrastar alguém para si mesmo, não."""
    final = {**responsavel_de, **mudancas}
    for chapa in mudancas:
        if final.get(chapa) == chapa:
            if aceita_topo:
                continue
            raise OrganogramaInvalido("Ninguém pode responder a si mesmo.")
        vistos = {chapa}
        acima = final.get(chapa)
        while acima:
            if acima in vistos:
                raise OrganogramaInvalido(
                    "Esse movimento faria a cadeia de responsáveis voltar para a mesma pessoa."
                )
            vistos.add(acima)
            proximo = final.get(acima)
            acima = proximo if proximo != acima else None


@bp.route("/api/indiretos/organograma", methods=["POST"])
@auth.exige_administrador
def mover_no_organograma():
    """Arrastar no organograma: para cada pessoa, a quem ela responde (o mesmo
    responsável dos cartões), se é assessoria dele e, opcionalmente, abaixo
    de qual bloco de função da equipe dele aparece. Tudo ou nada."""
    mudancas = _corpo().get("mudancas")
    if not isinstance(mudancas, list) or not mudancas:
        return jsonify({"erro": "Nenhuma mudança enviada."}), 400

    session = SessionLocal()
    try:
        alocacoes = {a.CHAPA: a for a in session.query(AlocacaoIndireto).all()}
        novos = {}
        for mudanca in mudancas:
            chapa = chapa_da_celula((mudanca or {}).get("chapa"))
            responsavel = chapa_da_celula((mudanca or {}).get("responsavel_chapa"))
            if chapa not in alocacoes:
                return jsonify({"erro": f"Indireto {chapa} não está alocado."}), 404
            if not responsavel or not _colaborador(session, responsavel):
                return jsonify({"erro": "Responsável não encontrado no cadastro."}), 404
            abaixo = str(mudanca.get("abaixo_de_funcao") or "").strip() or None
            novos[chapa] = (responsavel, bool(mudanca.get("assessoria")), abaixo)

        try:
            validar_cadeia(
                {c: a.RESPONSAVEL_CHAPA for c, a in alocacoes.items() if a.RESPONSAVEL_CHAPA},
                {c: r for c, (r, _, _) in novos.items()},
            )
        except OrganogramaInvalido as erro:
            return jsonify({"erro": str(erro)}), 400

        for chapa, (responsavel, assessoria, abaixo) in novos.items():
            alocacao = alocacoes[chapa]
            alocacao.RESPONSAVEL_CHAPA = responsavel
            alocacao.ASSESSORIA = assessoria
            alocacao.ABAIXO_DE_FUNCAO = abaixo
            alocacao.ALOCADO_POR = quem_esta_logado()
            alocacao.ALOCADO_EM = datetime.now(timezone.utc)
        session.commit()
        return jsonify({"sucesso": True})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] mover_no_organograma: {erro}")
        return jsonify({"erro": "Não foi possível mover no organograma."}), 500
    finally:
        session.close()


def plano_bases_responsavel(alocacoes, chapa, bases):
    """Responsável por base. alocacoes: {chapa: AlocacaoIndireto}. Devolve
    (novos_responsaveis {chapa: responsavel}, bases_por_chapa {chapa: [bases]}).

    - base marcada: as pessoas do setor naquela base que respondiam ao chefe
      dele (ou a quem respondia pela base antes) passam a responder a ele, e
      a base sai de quem respondia por ela;
    - base desmarcada: quem daquela base respondia a ele volta para o chefe
      dele."""
    pessoa = alocacoes[chapa]
    setor = pessoa.SETOR
    chefe = pessoa.RESPONSAVEL_CHAPA if pessoa.RESPONSAVEL_CHAPA != chapa else None
    antigas = set(bases_da_coluna(pessoa.BASES_RESPONSAVEL))
    novas = set(bases)
    do_setor = [a for a in alocacoes.values() if a.SETOR == setor]

    responsaveis = {}
    bases_por_chapa = {chapa: sorted(novas)}
    for base in novas - antigas:
        anteriores = {
            a.CHAPA for a in do_setor
            if a.CHAPA != chapa and base in bases_da_coluna(a.BASES_RESPONSAVEL)
        }
        for c in anteriores:
            restantes = [b for b in bases_da_coluna(alocacoes[c].BASES_RESPONSAVEL) if b != base]
            bases_por_chapa[c] = restantes
        if setor == SETOR_GERENCIA:
            continue
        para_mim = {chefe} | anteriores if chefe else anteriores
        for a in do_setor:
            if a.BASE == base and a.CHAPA not in (chapa, chefe) and a.RESPONSAVEL_CHAPA in para_mim:
                responsaveis[a.CHAPA] = chapa
    for base in antigas - novas:
        if not chefe or setor == SETOR_GERENCIA:
            continue
        for a in do_setor:
            if a.BASE == base and a.RESPONSAVEL_CHAPA == chapa:
                responsaveis[a.CHAPA] = chefe
    return responsaveis, bases_por_chapa


@bp.route("/api/indiretos/responsavel-base", methods=["POST"])
@auth.exige_administrador
def salvar_responsavel_base():
    """Marca as bases do setor pelas quais a pessoa responde (ver
    plano_bases_responsavel)."""
    dados = _corpo()
    chapa = chapa_da_celula(dados.get("chapa"))
    bases = sorted({texto_padrao(b) for b in (dados.get("bases") or []) if texto_padrao(b)})

    session = SessionLocal()
    try:
        alocacoes = {a.CHAPA: a for a in session.query(AlocacaoIndireto).all()}
        if chapa not in alocacoes:
            return jsonify({"erro": "Aloque a pessoa num setor antes de torná-la responsável por base."}), 404
        responsaveis, bases_por_chapa = plano_bases_responsavel(alocacoes, chapa, bases)
        try:
            validar_cadeia(
                {c: a.RESPONSAVEL_CHAPA for c, a in alocacoes.items() if a.RESPONSAVEL_CHAPA},
                responsaveis,
            )
        except OrganogramaInvalido as erro:
            return jsonify({"erro": str(erro)}), 400

        agora = datetime.now(timezone.utc)
        quem = quem_esta_logado()
        for c, responsavel in responsaveis.items():
            alocacoes[c].RESPONSAVEL_CHAPA = responsavel
            alocacoes[c].ASSESSORIA = False
            alocacoes[c].ABAIXO_DE_FUNCAO = None
            alocacoes[c].ALOCADO_POR, alocacoes[c].ALOCADO_EM = quem, agora
        for c, lista in bases_por_chapa.items():
            alocacoes[c].BASES_RESPONSAVEL = ",".join(lista) or None
        session.commit()
        return jsonify({"sucesso": True, "movidos": len(responsaveis)})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] salvar_responsavel_base: {erro}")
        return jsonify({"erro": "Não foi possível salvar o responsável pela base."}), 500
    finally:
        session.close()


@bp.route("/api/indiretos/setor-gerente", methods=["POST"])
@auth.exige_administrador
def ligar_setor_ao_gerente():
    """Arrastar um setor no organograma geral: o setor inteiro (todas as bases)
    fica abaixo do gerente ou, com apoio, ao lado da linha. Com "responsaveis",
    vale só para esses responsáveis do setor e quem está abaixo deles. Sem
    gerente: desfaz a ligação, e volta a seguir o resto do setor."""
    dados = _corpo()
    setor = texto_padrao(dados.get("setor"))
    gerente = chapa_da_celula(dados.get("gerente_chapa"))
    responsaveis = [chapa_da_celula(c) for c in (dados.get("responsaveis") or [])] or [""]
    if not setor:
        return jsonify({"erro": "Informe o setor."}), 400

    session = SessionLocal()
    try:
        if gerente:
            alocacao = session.query(AlocacaoIndireto).filter(AlocacaoIndireto.CHAPA == gerente).first()
            if not alocacao or alocacao.SETOR != SETOR_GERENCIA:
                return jsonify({"erro": "Essa pessoa não está na Gerência."}), 404
        for responsavel in responsaveis:
            ligacao = session.get(SetorGerencia, (setor, responsavel))
            if not gerente:
                if ligacao:
                    session.delete(ligacao)
                continue
            if not ligacao:
                ligacao = SetorGerencia(SETOR=setor, RESPONSAVEL_CHAPA=responsavel)
                session.add(ligacao)
            ligacao.GERENTE_CHAPA = gerente
            ligacao.APOIO = bool(dados.get("apoio"))
            ligacao.ALOCADO_POR = quem_esta_logado()
            ligacao.ALOCADO_EM = datetime.now(timezone.utc)
        session.commit()
        return jsonify({"sucesso": True})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] ligar_setor_ao_gerente: {erro}")
        return jsonify({"erro": "Não foi possível ligar o setor ao gerente."}), 500
    finally:
        session.close()


@bp.route("/api/indiretos/alocacao/<chapa>", methods=["DELETE"])
@auth.exige_administrador
def remover_alocacao_indireto(chapa):
    session = SessionLocal()
    try:
        removidas = (
            session.query(AlocacaoIndireto)
            .filter(AlocacaoIndireto.CHAPA == chapa_da_celula(chapa))
            .delete(synchronize_session=False)
        )
        session.commit()
        if not removidas:
            return jsonify({"erro": "Este indireto não está alocado."}), 404
        return jsonify({"sucesso": True})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] remover_alocacao_indireto: {erro}")
        return jsonify({"erro": "Não foi possível tirar do setor."}), 500
    finally:
        session.close()


# ============================================================
# PLANILHA: GERAR
# ============================================================

def montar_planilha_indiretos(session):
    """Planilha de cadastro já com as alocações de hoje, mais as abas de apoio
    (Indiretos, com quem falta; Colaboradores, para achar o responsável;
    Bases). As colunas de nome são fórmulas só para conferência: a análise
    lê apenas SETOR, BASE, as duas chapas, a observação e o subsetor."""
    colaboradores = session.query(Colaborador).order_by(Colaborador.NOME).all()
    alocacoes = (
        session.query(AlocacaoIndireto)
        .order_by(AlocacaoIndireto.SETOR, AlocacaoIndireto.BASE, AlocacaoIndireto.CHAPA)
        .all()
    )
    indiretos = [c for c in colaboradores if (c.TIPO_FUNÇÃO or "").upper() == TIPO_INDIRETO]
    bases = list(BASES_PADRAO) + sorted({a.BASE for a in alocacoes} - set(BASES_PADRAO))

    livro = Workbook()
    aba = livro.active
    aba.title = ABA_ALOCACAO
    for i, titulo in enumerate(COLUNAS_PLANILHA, start=1):
        aba.cell(row=1, column=i, value=titulo)

    preencher = {1, 2, 3, 5, 8, 9}
    ultima = len(alocacoes) + LINHAS_EXTRAS_PLANILHA + 1

    def busca(celula, coluna):
        return (
            f'=IF({celula}="","",IFERROR(INDEX(Colaboradores!${coluna}:${coluna},'
            f'MATCH({celula}&"",Colaboradores!$A:$A,0)),"CHAPA NÃO ENCONTRADA"))'
        )

    for n, alocacao in enumerate(alocacoes, start=2):
        aba.cell(row=n, column=1, value=alocacao.SETOR)
        aba.cell(row=n, column=2, value=alocacao.BASE)
        aba.cell(row=n, column=3, value=alocacao.RESPONSAVEL_CHAPA or "")
        aba.cell(row=n, column=5, value=alocacao.CHAPA)
        aba.cell(row=n, column=8, value=alocacao.OBSERVACAO or "")
        aba.cell(row=n, column=9, value=alocacao.SUBSETOR or "")
    # só o que veio do banco: as colunas de busca logo abaixo são fórmula de verdade
    estilo.gravar_formulas_como_texto(aba)

    for n in range(2, ultima + 1):
        for c in range(1, len(COLUNAS_PLANILHA) + 1):
            celula = aba.cell(row=n, column=c)
            celula.border = estilo.BORDA
            celula.font = Font(name=estilo.FONTE, size=10, color="0000FF" if c in preencher else "000000")
            celula.fill = estilo.FUNDO_PREENCHER if c in preencher else estilo.FUNDO_CONSULTA
        aba.cell(row=n, column=3).number_format = "@"
        aba.cell(row=n, column=5).number_format = "@"
        aba.cell(row=n, column=4, value=busca(f"C{n}", "B"))
        aba.cell(row=n, column=6, value=busca(f"E{n}", "B"))
        aba.cell(row=n, column=7, value=busca(f"E{n}", "C"))

    for c, (titulo, largura) in enumerate(zip(COLUNAS_PLANILHA, (26, 12, 18, 34, 16, 34, 30, 30, 18)), start=1):
        celula = aba.cell(row=1, column=c)
        celula.font, celula.fill, celula.border = estilo.FONTE_CABECALHO, estilo.FUNDO_CABECALHO, estilo.BORDA
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.comment = Comment(
            "Preencha esta coluna." if c in preencher
            else "Automático: confere a chapa no cadastro. Não digite aqui.",
            "Gerenciador de Equipes",
        )
        aba.column_dimensions[celula.column_letter].width = largura
    aba.row_dimensions[1].height = 30
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = f"A1:I{ultima}"
    aba.sheet_properties.tabColor = estilo.BORDO

    lista_bases = DataValidation(
        type="list", formula1=f"=Bases!$A$2:$A${len(bases) + 1}", allow_blank=True,
        errorStyle="warning", errorTitle="Base fora da lista",
        error="Essa sigla não está na lista de bases. Confirme se está certa.",
    )
    aba.add_data_validation(lista_bases)
    lista_bases.add(f"B2:B{ultima}")

    vermelho = estilo.fundo_condicional("FBE9E7")
    fonte_vermelha = Font(name=estilo.FONTE, size=10, bold=True, color="B3261E")
    for faixa, ref in ((f"D2:D{ultima}", "D2"), (f"F2:G{ultima}", "F2")):
        aba.conditional_formatting.add(
            faixa, FormulaRule(formula=[f'{ref}="CHAPA NÃO ENCONTRADA"'], fill=vermelho, font=fonte_vermelha)
        )
    aba.conditional_formatting.add(
        f"E2:E{ultima}",
        FormulaRule(
            formula=[f'AND(E2<>"",COUNTIF($E$2:$E${ultima},E2)>1)'],
            fill=estilo.fundo_condicional("FDF0CD"), font=estilo.FONTE_ALERTA,
        ),
    )

    # ---------- Indiretos (com quem já está na planilha)
    aba_ind = livro.create_sheet("Indiretos")
    cab = ["CHAPA", "NOME", "FUNÇÃO", "BASE (CADASTRO)", "SITUAÇÃO", "JÁ NA PLANILHA?"]
    for i, t in enumerate(cab, start=1):
        aba_ind.cell(row=1, column=i, value=t)
    for n, c in enumerate(indiretos, start=2):
        valores = [c.CHAPA, c.NOME or "", c.FUNÇÃO or "", base_do_cadastro(c), c.SITUAÇÃO or ""]
        for i, v in enumerate(valores, start=1):
            aba_ind.cell(row=n, column=i, value=v)
        aba_ind.cell(row=n, column=6, value=f'=IF(COUNTIF({ABA_ALOCACAO}!$E$2:$E${ultima},A{n}&"")>0,"SIM","FALTA")')
    estilo.estilizar_aba(aba_ind, consulta=True, colunas_com_formula=("JÁ NA PLANILHA?",))
    fim = len(indiretos) + 1
    aba_ind.conditional_formatting.add(f"F2:F{fim}", FormulaRule(
        formula=['F2="FALTA"'], fill=estilo.fundo_condicional("FDF0CD"), font=estilo.FONTE_ALERTA))
    aba_ind.conditional_formatting.add(f"F2:F{fim}", FormulaRule(
        formula=['F2="SIM"'], fill=estilo.fundo_condicional("E6F2EA"),
        font=Font(name=estilo.FONTE, size=10, bold=True, color="1B6B45")))
    aba_ind["H1"], aba_ind["H2"], aba_ind["H3"] = "Indiretos no cadastro", "Já na planilha", "Faltam"
    aba_ind["I1"] = f"=COUNTA(A2:A{fim})"
    aba_ind["I2"] = f'=COUNTIF(F2:F{fim},"SIM")'
    aba_ind["I3"] = "=I1-I2"
    for ref in ("H1", "H2", "H3"):
        aba_ind[ref].font = Font(name=estilo.FONTE, size=10, bold=True, color=estilo.BORDO)
    for ref in ("I1", "I2", "I3"):
        aba_ind[ref].font = Font(name=estilo.FONTE, size=10, bold=True)
    aba_ind.column_dimensions["H"].width = 22

    # ---------- Colaboradores (para achar o responsável) e Bases
    aba_col = livro.create_sheet("Colaboradores")
    for i, t in enumerate(["CHAPA", "NOME", "FUNÇÃO", "TIPO", "BASE (CADASTRO)", "SITUAÇÃO"], start=1):
        aba_col.cell(row=1, column=i, value=t)
    for n, c in enumerate(colaboradores, start=2):
        for i, v in enumerate([c.CHAPA, c.NOME or "", c.FUNÇÃO or "", c.TIPO_FUNÇÃO or "",
                               base_do_cadastro(c), c.SITUAÇÃO or ""], start=1):
            aba_col.cell(row=n, column=i, value=v)
    estilo.estilizar_aba(aba_col, consulta=True)

    aba_bases = livro.create_sheet("Bases")
    aba_bases.cell(row=1, column=1, value="BASE")
    for n, b in enumerate(bases, start=2):
        aba_bases.cell(row=n, column=1, value=b)
    estilo.estilizar_aba(aba_bases, consulta=True)

    estilo.adicionar_instrucoes(
        livro,
        "Planilha de alocação de indiretos",
        [
            ("Para que serve", "Informar em qual SETOR e BASE cada indireto trabalha e quem é o RESPONSÁVEL dele. Uma linha por indireto. O arquivo já vem com as alocações de hoje."),
            ("Onde preencher", "Aba ALOCAÇÃO, nas colunas amarelas: SETOR, BASE, CHAPA RESPONSÁVEL, CHAPA INDIRETO, OBSERVAÇÃO (opcional) e SUBSETOR (opcional)."),
            ("Colunas automáticas", "NOME RESPONSÁVEL, NOME INDIRETO e FUNÇÃO INDIRETO (cinza) se preenchem pela chapa, só para conferir. \"CHAPA NÃO ENCONTRADA\" em vermelho = chapa errada ou fora do cadastro."),
            ("SETOR", "Escreva sempre do mesmo jeito (o sistema guarda em maiúsculas). O mesmo setor em várias bases aparece na tela num cartão só, com uma seção por base."),
            ("BASE", "Escolha na lista. Você decide a base, independente da seção do cadastro da pessoa."),
            ("RESPONSÁVEL", "Chapa de quem responde pelo indireto. É por pessoa: no mesmo setor e base pode haver responsáveis diferentes. Pode ser qualquer colaborador (aba COLABORADORES)."),
            ("CHAPA INDIRETO", "Cada indireto uma vez só (repetido fica laranja). A aba INDIRETOS mostra quem ainda falta."),
            ("Nunca tira ninguém", "Quem não estiver na planilha continua onde está. Para tirar alguém do setor, use o menu da pessoa na tela Indiretos."),
            ("Antes de aplicar", "Envie o arquivo em Indiretos > Planilha de cadastro: o sistema mostra uma prévia do que muda e só grava quando você confirmar."),
        ],
    )
    livro.active = 0
    return livro


@bp.route("/api/indiretos/planilha", methods=["GET"])
@auth.exige_administrador
def baixar_planilha_indiretos():
    session = SessionLocal()
    try:
        livro = montar_planilha_indiretos(session)
        buffer = io.BytesIO()
        livro.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="alocacao-indiretos.xlsx",
        )
    except Exception as erro:
        print(f"[ERRO] baixar_planilha_indiretos: {erro}")
        return jsonify({"erro": "Não foi possível gerar a planilha."}), 500
    finally:
        session.close()


# ============================================================
# PLANILHA: ANALISAR / APLICAR
# ============================================================

def _ler_aba(arquivo):
    """Aba "Alocação" pelo nome (o modelo antigo tinha Instruções primeiro);
    se não houver, a primeira aba que tenha as colunas."""
    try:
        abas = pd.read_excel(arquivo, sheet_name=None, dtype=object)
    except Exception as erro:
        raise ValueError(f"Não foi possível ler a planilha: {erro}")
    df = abas.get(ABA_ALOCACAO)
    if df is None:
        df = next(
            (d for d in abas.values() if set(COLUNAS_OBRIGATORIAS) <= {str(c).strip() for c in d.columns}),
            None,
        )
    if df is None:
        raise ValueError(f"A planilha precisa da aba \"{ABA_ALOCACAO}\" com as colunas {', '.join(COLUNAS_OBRIGATORIAS)}.")
    df.columns = [str(c).strip() for c in df.columns]
    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in df.columns]
    if faltando:
        raise ValueError(f"A planilha precisa das colunas {', '.join(faltando)}.")
    return df


def analisar_planilha_indiretos(arquivo, session):
    """Monta o plano sem gravar nada: alocar (novos), mover (setor ou base
    mudou), atualizar (só responsável/observação/subsetor mudou), sem mudança e erros."""
    df = _ler_aba(arquivo)
    colaboradores = {c.CHAPA: c for c in session.query(Colaborador).all()}
    atuais = {a.CHAPA: a for a in session.query(AlocacaoIndireto).all()}

    plano = {"alocar": [], "mover": [], "atualizar": [], "sem_mudanca": 0, "erros": []}
    vistos = {}

    for indice, linha in df.iterrows():
        numero = int(indice) + 2
        setor = texto_padrao(linha.get(COL_SETOR))
        base = texto_padrao(linha.get(COL_BASE))
        responsavel = chapa_da_celula(linha.get(COL_RESPONSAVEL))
        chapa = chapa_da_celula(linha.get(COL_INDIRETO))
        observacao = texto_livre(linha.get(COL_OBSERVACAO)) if COL_OBSERVACAO in df.columns else None
        # coluna opcional: planilha sem ela não mexe no subsetor de ninguém
        tem_subsetor = COL_SUBSETOR in df.columns
        subsetor = (texto_padrao(linha.get(COL_SUBSETOR)) or None) if tem_subsetor else None

        if not (setor or base or responsavel or chapa):
            continue  # linha preparada e não usada

        def erro(mensagem):
            plano["erros"].append({"linha": numero, "chapa": chapa or "—", "erro": mensagem})

        faltam = [n for n, v in (("SETOR", setor), ("BASE", base),
                                 ("CHAPA RESPONSÁVEL", responsavel), ("CHAPA INDIRETO", chapa)) if not v]
        if faltam:
            erro(f"Falta preencher: {', '.join(faltam)}.")
            continue
        pessoa = colaboradores.get(chapa)
        if not pessoa:
            erro("CHAPA INDIRETO não encontrada no cadastro.")
            continue
        if (pessoa.TIPO_FUNÇÃO or "").upper() != TIPO_INDIRETO:
            erro(f"{pessoa.NOME} está como {pessoa.TIPO_FUNÇÃO or 'sem tipo'} no cadastro, não INDIRETO.")
            continue
        chefe = colaboradores.get(responsavel)
        if not chefe:
            erro("CHAPA RESPONSÁVEL não encontrada no cadastro.")
            continue
        if chapa in vistos:
            erro(f"Indireto repetido: já aparece na linha {vistos[chapa]}.")
            continue
        vistos[chapa] = numero

        item = {
            "linha": numero, "chapa": chapa, "nome": pessoa.NOME,
            "setor": setor, "base": base,
            "responsavel_chapa": responsavel, "responsavel_nome": chefe.NOME,
            "observacao": observacao,
        }
        atual = atuais.get(chapa)
        if not tem_subsetor:
            subsetor = (atual.SUBSETOR if atual else None) or None
        item["subsetor"] = subsetor
        if not atual:
            plano["alocar"].append(item)
        elif (atual.SETOR, atual.BASE) != (setor, base):
            item["de"] = f"{atual.SETOR} · {atual.BASE}"
            plano["mover"].append(item)
        elif (
            atual.RESPONSAVEL_CHAPA != responsavel
            or (atual.OBSERVACAO or None) != observacao
            or (atual.SUBSETOR or None) != subsetor
        ):
            mudou = []
            if atual.RESPONSAVEL_CHAPA != responsavel:
                antigo = colaboradores.get(atual.RESPONSAVEL_CHAPA)
                mudou.append(f"responsável: {antigo.NOME if antigo else atual.RESPONSAVEL_CHAPA or '—'} → {chefe.NOME}")
            if (atual.OBSERVACAO or None) != observacao:
                mudou.append("observação")
            if (atual.SUBSETOR or None) != subsetor:
                mudou.append(f"subsetor: {atual.SUBSETOR or '—'} → {subsetor or '—'}")
            item["mudancas"] = mudou
            plano["atualizar"].append(item)
        else:
            plano["sem_mudanca"] += 1

    # a cadeia só dá para conferir com todas as linhas aplicadas: A -> B numa
    # linha e B -> A em outra só formam ciclo juntas
    mudancas = plano["alocar"] + plano["mover"] + plano["atualizar"]
    cadeia_final = {
        **{c: a.RESPONSAVEL_CHAPA for c, a in atuais.items()},
        **{item["chapa"]: item["responsavel_chapa"] for item in mudancas},
    }
    for item in mudancas:
        try:
            validar_cadeia(cadeia_final, {item["chapa"]: item["responsavel_chapa"]}, aceita_topo=True)
        except OrganogramaInvalido as erro:
            plano["erros"].append({"linha": item["linha"], "chapa": item["chapa"], "erro": str(erro)})

    return plano


def _arquivo_enviado():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        raise ValueError("Arquivo não enviado.")
    return arquivo


@bp.route("/api/indiretos/planilha/previa", methods=["POST"])
@auth.exige_administrador
def prever_planilha_indiretos():
    session = SessionLocal()
    try:
        return jsonify(analisar_planilha_indiretos(_arquivo_enviado(), session))
    except ValueError as erro:
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        print(f"[ERRO] prever_planilha_indiretos: {erro}")
        return jsonify({"erro": "Não foi possível analisar a planilha."}), 500
    finally:
        session.close()


@bp.route("/api/indiretos/planilha/aplicar", methods=["POST"])
@auth.exige_administrador
def aplicar_planilha_indiretos():
    """Reanalisa (o banco pode ter mudado desde a prévia) e grava. Com
    qualquer erro na planilha, não grava nada."""
    session = SessionLocal()
    try:
        plano = analisar_planilha_indiretos(_arquivo_enviado(), session)
        if plano["erros"]:
            return jsonify({"erro": "Corrija os erros da planilha antes de aplicar.", **plano}), 400

        agora = datetime.now(timezone.utc)
        quem = quem_esta_logado()
        atuais = {a.CHAPA: a for a in session.query(AlocacaoIndireto).all()}
        for item in plano["alocar"] + plano["mover"] + plano["atualizar"]:
            alocacao = atuais.get(item["chapa"])
            if not alocacao:
                alocacao = AlocacaoIndireto(CHAPA=item["chapa"])
                session.add(alocacao)
            alocacao.SETOR = item["setor"]
            alocacao.BASE = item["base"]
            if alocacao.RESPONSAVEL_CHAPA != item["responsavel_chapa"]:
                alocacao.ASSESSORIA = False
            alocacao.RESPONSAVEL_CHAPA = item["responsavel_chapa"]
            alocacao.OBSERVACAO = item["observacao"]
            alocacao.SUBSETOR = item["subsetor"]
            alocacao.ALOCADO_POR = quem
            alocacao.ALOCADO_EM = agora
        session.commit()
        return jsonify({
            "sucesso": True,
            "alocados": len(plano["alocar"]),
            "movidos": len(plano["mover"]),
            "atualizados": len(plano["atualizar"]),
            "sem_mudanca": plano["sem_mudanca"],
        })
    except ValueError as erro:
        session.rollback()
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] aplicar_planilha_indiretos: {erro}")
        return jsonify({"erro": "Não foi possível aplicar a planilha."}), 500
    finally:
        session.close()
